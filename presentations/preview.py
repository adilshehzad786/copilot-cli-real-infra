#!/usr/bin/env python3
"""Render each slide of the built deck to a PNG, and flag text that overflows.

    python3 -m pip install python-pptx pillow
    python3 presentations/preview.py

This reads the generated .pptx rather than the content module, so it checks what
is actually in the file that gets presented. It is a legibility check, not a
faithful PowerPoint renderer: it draws the rectangles and the text runs at their
real positions, sizes and colours, which is enough to catch the two failures that
matter — a line that runs past its box, and a slide that is too dense to read
from the back of a room.

Output goes to presentations/previews/ (git-ignored).
"""
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Emu

SCALE = 2  # 960x540pt -> 1920x1080px
FONT_DIR = Path("/System/Library/Fonts/Supplemental")
FALLBACK = Path("/Library/Fonts")
FACES = {
    ("Arial", False): "Arial.ttf",
    ("Arial", True): "Arial Bold.ttf",
    ("Courier New", False): "Courier New.ttf",
    ("Courier New", True): "Courier New Bold.ttf",
}
_cache = {}


def font(name, bold, size_pt):
    key = (name, bold, size_pt)
    if key not in _cache:
        filename = FACES.get((name, bold), FACES[("Arial", bold)])
        for directory in (FONT_DIR, FALLBACK):
            candidate = directory / filename
            if candidate.is_file():
                _cache[key] = ImageFont.truetype(str(candidate), int(size_pt * SCALE))
                break
        else:
            _cache[key] = ImageFont.load_default()
    return _cache[key]


def pt(emu):
    return Emu(emu).pt * SCALE


def render(slide, index, out_dir):
    image = Image.new("RGB", (960 * SCALE, 540 * SCALE), (0x0D, 0x11, 0x17))
    draw = ImageDraw.Draw(image)
    warnings = []
    for shape in slide.shapes:
        x, y = pt(shape.left), pt(shape.top)
        w, h = pt(shape.width), pt(shape.height)
        if shape.shape_type is not None and not shape.has_text_frame:
            continue
        fill = getattr(shape, "fill", None)
        try:
            if fill is not None and fill.type is not None and fill.type == 1:
                draw.rectangle([x, y, x + w, y + h], fill=tuple(fill.fore_color.rgb))
        except (AttributeError, TypeError, ValueError):
            pass
        if not shape.has_text_frame:
            continue
        frame = shape.text_frame
        if not any(r.text.strip() for para in frame.paragraphs for r in para.runs):
            continue
        wrap = frame.word_wrap is not False
        lines = []          # (runs, size, spacing, alignment)
        for para in frame.paragraphs:
            runs = [(r.text, r.font) for r in para.runs]
            spacing = para.line_spacing or 1.15
            align = str(para.alignment)
            if not runs:
                lines.append(([], 14, spacing, align))
                continue
            size = max((r[1].size.pt if r[1].size else 14) for r in runs)
            if not wrap:
                lines.append((runs, size, spacing, align))
                continue
            # Greedy wrap, preserving run styling across the break.
            current, used = [], 0.0
            for value, f in runs:
                face = font(f.name or "Arial", bool(f.bold), f.size.pt if f.size else 14)
                for word in value.split(" "):
                    piece = word if not current and used == 0 else " " + word
                    width = draw.textlength(piece, font=face)
                    if used + width > w and current:
                        lines.append((current, size, spacing, align))
                        current, used = [(word, f)], draw.textlength(word, font=face)
                    else:
                        current.append((piece, f))
                        used += width
            if current:
                lines.append((current, size, spacing, align))
        total = sum(size * spacing * SCALE for _, size, spacing, _ in lines)
        anchor = str(frame.vertical_anchor)
        cursor = y + (0 if "TOP" in anchor else (h - total) / 2)
        for runs, size, spacing, align in lines:
            line_h = size * spacing * SCALE
            if not runs:
                cursor += line_h
                continue
            widths = [draw.textlength(v, font=font(
                f.name or "Arial", bool(f.bold), f.size.pt if f.size else 14))
                for v, f in runs]
            start = x + ((w - sum(widths)) / 2 if "CENTER" in align else 0)
            cur = start
            for (value, f), width in zip(runs, widths):
                colour = (0xE6, 0xED, 0xF3)
                try:
                    colour = tuple(f.color.rgb)
                except (AttributeError, TypeError, ValueError):
                    pass
                draw.text((cur, cursor), value, fill=colour, font=font(
                    f.name or "Arial", bool(f.bold), f.size.pt if f.size else 14))
                cur += width
            cursor += line_h
        if cursor > y + h + 3 * SCALE:
            warnings.append(f"  text overruns its box by {(cursor - y - h)/SCALE:.0f}pt "
                            f"(box y={y/SCALE:.0f} h={h/SCALE:.0f}): "
                            f"{lines[0][0][0][0][:50]!r}")
        if cursor > 540 * SCALE:
            warnings.append(f"  text runs off the bottom of the slide")
    out_dir.mkdir(parents=True, exist_ok=True)
    image.save(out_dir / f"slide-{index:02d}.png")
    return warnings


def main():
    here = Path(__file__).resolve().parent
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else \
        here / "copilot-cli-real-infra-deck.pptx"
    out_dir = here / "previews"
    prs = Presentation(str(source))
    problems = 0
    for index, slide in enumerate(prs.slides, 1):
        warnings = render(slide, index, out_dir)
        if warnings:
            problems += len(warnings)
            print(f"slide {index}:")
            for line in warnings:
                print(line)
    print(f"\nRendered {len(prs.slides._sldIdLst)} slides to {out_dir}")
    print("No overflow detected." if not problems else f"{problems} overflow warning(s).")
    return 1 if problems else 0


if __name__ == "__main__":
    sys.exit(main())
