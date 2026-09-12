"""Design system for the talk deck: GitHub-dark terminal, 960x540pt canvas.

Layout primitives only. Slide copy lives in ``deck_content.py`` and the build
entry point is ``build_deck.py``. Keeping the three apart means a copy change
never risks the geometry, and a geometry change is reviewable on its own.
"""
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

# GitHub dark palette, taken from the original deck so the visual identity survives.
CANVAS = RGBColor(0x0D, 0x11, 0x17)
DEEP = RGBColor(0x01, 0x04, 0x09)
SURFACE = RGBColor(0x16, 0x1B, 0x22)
RAISED = RGBColor(0x1C, 0x21, 0x28)
INK = RGBColor(0xE6, 0xED, 0xF3)
MUTED = RGBColor(0x8B, 0x94, 0x9E)
ACCENT = RGBColor(0xF0, 0x88, 0x3E)
GREEN = RGBColor(0x3F, 0xB9, 0x50)
BLUE = RGBColor(0x58, 0xA6, 0xFF)
RED = RGBColor(0xF8, 0x51, 0x49)
WARM = RGBColor(0x2B, 0x1D, 0x10)
LEAF = RGBColor(0x0D, 0x28, 0x18)

SANS = "Arial"
MONO = "Courier New"

# 13.333in x 7.5in at 72pt/in. Everything below is in points.
SLIDE_W, SLIDE_H = 960.0, 540.0
MARGIN = 47.0
COL = SLIDE_W - 2 * MARGIN  # 866pt of usable width
TITLE_Y = 30.0
BODY_TOP = 112.0
KICKER_Y = 470.0

# Type scale. Raised from the original deck: nothing below 16pt on a slide body,
# because 13pt on this canvas is roughly 26px on a 1080p projector and the back
# of a conference room cannot read it.
SZ_TITLE_HERO = 46
SZ_STATEMENT = 38
SZ_TITLE = 32
SZ_LEAD = 20
SZ_BODY = 17
SZ_META = 16
SZ_KICKER = 17
SZ_NUM = 20


def new_slide(prs):
    """A blank slide painted with the canvas colour."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = CANVAS
    return slide


def rect(slide, x, y, w, h, fill=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    shape.line.fill.background()
    shape.shadow.inherit = False
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    return shape


def text(slide, x, y, w, h, runs, size=SZ_BODY, color=INK, bold=False,
         font=SANS, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, spacing=1.15):
    """Draw one text box. ``runs`` is a string, or a list of strings (one per
    line), or a list of (text, colour, font) tuples for mixed styling."""
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = anchor
    lines = [runs] if isinstance(runs, str) else list(runs)
    for index, line in enumerate(lines):
        para = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        para.alignment = align
        para.line_spacing = spacing
        for piece in (line if isinstance(line, list) else [line]):
            value, tint, face = piece if isinstance(piece, tuple) else (piece, color, font)
            run = para.add_run()
            run.text = value
            run.font.size = Pt(size)
            run.font.color.rgb = tint
            run.font.bold = bold
            run.font.name = face
    return box


def prompt_badge(slide, y=37.0):
    """The orange `$` chip that marks a content slide."""
    rect(slide, MARGIN, y, 26, 26, ACCENT)
    text(slide, MARGIN, y, 26, 26, "$", size=17, color=CANVAS, bold=True,
         font=MONO, align=PP_ALIGN.CENTER)


def title(slide, heading, y=TITLE_Y):
    prompt_badge(slide, y + 7)
    text(slide, MARGIN + 40, y, COL - 40, 46, heading, size=SZ_TITLE, bold=True)


def kicker(slide, line, color=ACCENT, y=KICKER_Y):
    """The one-sentence takeaway pinned to the bottom of a content slide."""
    text(slide, MARGIN, y, COL, 42, line, size=SZ_KICKER, color=color, bold=True)


def rows(slide, entries, top=BODY_TOP, height=54.0, gap=8.0, label_w=520.0):
    """Numbered cards: (number, label, note, highlight)."""
    for index, entry in enumerate(entries):
        number, label, note = entry[0], entry[1], entry[2]
        highlight = entry[3] if len(entry) > 3 else False
        y = top + index * (height + gap)
        rect(slide, MARGIN, y, COL, height, WARM if highlight else SURFACE)
        text(slide, MARGIN + 18, y, 40, height, str(number), size=SZ_NUM,
             color=ACCENT if highlight else MUTED, bold=True, font=MONO)
        text(slide, MARGIN + 64, y, label_w, height, label, size=SZ_BODY,
             color=INK, bold=highlight)
        text(slide, MARGIN + 64 + label_w + 16, y,
             COL - label_w - 96, height, note, size=SZ_META,
             color=ACCENT if highlight else MUTED)


def panels(slide, entries, top=BODY_TOP, height=250.0, gap=18.0, accent_top=True):
    """Side-by-side cards: (heading, [lines], colour)."""
    count = len(entries)
    width = (COL - gap * (count - 1)) / count
    for index, (heading, lines, tint) in enumerate(entries):
        x = MARGIN + index * (width + gap)
        rect(slide, x, top, width, height, SURFACE)
        if accent_top:
            rect(slide, x, top, width, 4, tint)
        text(slide, x + 22, top + 22, width - 44, 30, heading,
             size=SZ_LEAD, color=tint, bold=True, anchor=MSO_ANCHOR.TOP)
        text(slide, x + 22, top + 64, width - 44, height - 86, lines,
             size=SZ_BODY, color=INK, anchor=MSO_ANCHOR.TOP, spacing=1.35)


def terminal(slide, lines, x=MARGIN, y=BODY_TOP, w=COL, h=170.0):
    """A fake terminal block. ``lines`` are (text, colour) pairs."""
    rect(slide, x, y, w, h, DEEP)
    rect(slide, x, y, 4, h, ACCENT)
    body = [[(value, tint, MONO)] for value, tint in lines]
    text(slide, x + 26, y + 18, w - 52, h - 36, body, size=SZ_BODY,
         anchor=MSO_ANCHOR.TOP, spacing=1.45)


def callout(slide, label, body, tint=ACCENT, y=None, height=76.0):
    """The boxed aside the original deck used for gotchas and credit."""
    y = SLIDE_H - MARGIN - height if y is None else y
    background = WARM if tint is ACCENT else (LEAF if tint is GREEN else RAISED)
    rect(slide, MARGIN, y, COL, height, background)
    rect(slide, MARGIN, y, 4, height, tint)
    text(slide, MARGIN + 24, y + 12, COL - 48, 22, label, size=SZ_META,
         color=tint, bold=True, anchor=MSO_ANCHOR.TOP)
    text(slide, MARGIN + 24, y + 36, COL - 48, height - 48, body, size=SZ_BODY,
         color=INK, anchor=MSO_ANCHOR.TOP, spacing=1.3)


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body
